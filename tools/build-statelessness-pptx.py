#!/usr/bin/env python3
"""
build-statelessness-pptx.py — render the STATELESSNESS COMPENDIUM deck.

Companion to build-pptx.py (the main deck). It reuses every slide builder,
the theme, and the helpers from build-pptx.py — loaded via importlib because
that file's name has a hyphen — and supplies its own title / agenda / closing
slides plus the statelessness SECTIONS from sections_statelessness.py.

The existing deck build (build-pptx.py / sections.py) is untouched.

Diagrams: build-statelessness-deck.sh converts diagrams/statelessness/*.svg
into /tmp/diagrams-png/<name>.jpg before this runs.

Output: presentation/cpp-statelessness-compendium.pptx
"""
import importlib.util
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parent.parent
TOOLS = ROOT / "tools"


def load_build_pptx():
    """Load build-pptx.py (hyphenated filename) as a module."""
    spec = importlib.util.spec_from_file_location(
        "build_pptx", TOOLS / "build-pptx.py")
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


bp = load_build_pptx()
C, F, FontFam = bp.C, bp.F, bp.FontFam
SLIDE_W, SLIDE_H = bp.SLIDE_W, bp.SLIDE_H

OUT_PATH = ROOT / "presentation" / "cpp-statelessness-compendium.pptx"
NOTES_PATH = ROOT / "presentation" / "build-notes-statelessness.md"
REPO_URL = "github.com/patterncatalyst/cpp-container-optimization-tutorial"


def build_title(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bp.set_slide_bg(slide, C.BG_TITLE_DARK)

    dot_y = Inches(1.5)
    for i, color in enumerate([C.ACCENT_CYAN, C.ACCENT_GREEN, C.ACCENT_PURPLE]):
        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(1.0 + i * 0.35), dot_y, Inches(0.22), Inches(0.22))
        dot.fill.solid()
        dot.fill.fore_color.rgb = color
        dot.line.fill.background()

    bp.add_text(slide, Inches(1.0), Inches(2.0), Inches(11.3), Inches(0.5),
                "STATELESSNESS COMPENDIUM", size=Pt(20), color=C.ACCENT_CYAN,
                bold=True, font=FontFam.HEADER)
    bp.add_text(slide, Inches(1.0), Inches(2.5), Inches(11.3), Inches(1.5),
                "Stateless C++ Services\nUnder Container Orchestration",
                size=F.TITLE_HUGE, color=C.TEXT_WHITE, bold=True, font=FontFam.HEADER)
    bp.add_text(slide, Inches(1.0), Inches(4.5), Inches(11.3), Inches(0.5),
                "Request / process / deploy-time scope  •  RAII  •  PMR  •  pools  "
                "•  health & shutdown  •  the gRPC capstone",
                size=F.BODY, color=C.TEXT_LIGHT, font=FontFam.BODY)
    bp.add_text(slide, Inches(1.0), Inches(5.2), Inches(11.3), Inches(0.4),
                "Companion deck to \u201cOptimizing Modern C++ with Containers\u201d  "
                "\u2014  Docs 01\u201311",
                size=F.BODY_SMALL, color=C.TEXT_MUTED, italic=True, font=FontFam.BODY)
    bp.add_text(slide, Inches(1.0), Inches(5.9), Inches(11.3), Inches(0.4),
                "Every pattern has a runnable Podman demo  \u2014  02\u201311 host-verified",
                size=F.BODY_SMALL, color=C.ACCENT_CYAN, font=FontFam.BODY)
    bp.add_text(slide, Inches(1.0), Inches(6.5), Inches(11.3), Inches(0.4),
                REPO_URL, size=F.CAPTION, color=C.TEXT_MUTED, font=FontFam.CODE)

    bp.set_notes(slide,
        "This is the statelessness compendium deck — the architectural "
        "companion to the performance talk. Where the main deck is about "
        "making C++ fast under container constraints, this one is about "
        "making C++ services correct as stateless workloads under an "
        "orchestrator.\n\n"
        "The spine is a single idea developed across eleven documents: "
        "statelessness is a deployment posture, not a code property, and the "
        "C++ techniques that satisfy it — RAII for request scope, PMR for "
        "request memory, explicit process-scoped ownership, externalized "
        "authoritative state, health checks and graceful shutdown — compose "
        "into one realistic gRPC service by the end.\n\n"
        "Every pattern here has a runnable Podman demo, and examples 2 "
        "through 11 are host-verified end to end on Fedora 44. We'll cue each "
        "demo at the point its pattern is introduced.")


def build_agenda(prs, total, sections):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bp.set_slide_bg(slide, C.TEXT_WHITE)
    bp.add_header_bar(slide, "Overview", 2, total, "Agenda \u2014 the eleven documents")
    bp.add_footer(slide, "Agenda", 2, total)

    palette = [C.ACCENT_RED, C.ACCENT_ORANGE, C.ACCENT_BLUE, C.ACCENT_GREEN,
               C.ACCENT_CYAN, C.ACCENT_PURPLE]
    rows_per_col = 6
    col_width = Inches(6.0)
    row_height = Inches(0.62)
    grid_left = Inches(0.5)
    grid_top = Inches(1.1)

    for i, sec in enumerate(sections):
        color = palette[i % len(palette)]
        col = i // rows_per_col
        row = i % rows_per_col
        x = grid_left + col * (col_width + Inches(0.3))
        y = grid_top + row * row_height

        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, Inches(0.5), Inches(0.5))
        circle.fill.solid()
        circle.fill.fore_color.rgb = color
        circle.line.fill.background()
        bp.add_text(slide, x, y, Inches(0.5), Inches(0.5), str(sec["num"]),
                    size=Pt(16), color=C.TEXT_WHITE, bold=True,
                    align=PP_ALIGN.CENTER, anchor=bp.MSO_ANCHOR.MIDDLE,
                    font=FontFam.HEADER)
        bp.add_text(slide, x + Inches(0.65), y, Inches(5.2), Inches(0.5),
                    sec["title"], size=Pt(15), color=C.TEXT_DARK,
                    anchor=bp.MSO_ANCHOR.MIDDLE, font=FontFam.BODY)

    bp.set_notes(slide,
        "Eleven documents, three movements. Documents 1 through 3 establish "
        "the foundations — the scope vocabulary, RAII, and PMR. Documents 4 "
        "through 8 cover the operational concerns — process-scoped ownership, "
        "threading, 12-factor, state externalization, the ephemeral "
        "filesystem. Documents 9 through 11 land it — health checks and "
        "shutdown, the gRPC capstone that composes everything, and the build "
        "tooling appendix with the vendored helpers.\n\n"
        "Each section ends or pauses on a demo cue where the pattern has a "
        "runnable companion.")


def build_closing(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bp.set_slide_bg(slide, C.BG_TITLE_DARK)
    bp.add_text(slide, Inches(1.0), Inches(1.6), Inches(11.3), Inches(1.0),
                "The whole arc, runnable.", size=F.TITLE_LARGE, color=C.TEXT_WHITE,
                bold=True, font=FontFam.HEADER)
    bp.add_multi_text(slide, Inches(1.0), Inches(2.8), Inches(11.3), Inches(2.6), [
        dict(text="Statelessness is a deployment posture; the same binary earns it "
                  "by where it keeps state.", size=Pt(18), color=C.TEXT_LIGHT,
             font=FontFam.BODY),
        dict(text="Request scope \u2192 RAII + PMR. Process scope \u2192 owned in main(). "
                  "Authoritative state \u2192 externalized.", size=Pt(18),
             color=C.TEXT_LIGHT, font=FontFam.BODY, space_before=Pt(10)),
        dict(text="It all composes in the capstone \u2014 and examples 02\u201311 are "
                  "host-verified on Fedora 44.", size=Pt(18), color=C.ACCENT_CYAN,
             font=FontFam.BODY, space_before=Pt(10)),
    ])
    bp.add_text(slide, Inches(1.0), Inches(5.6), Inches(11.3), Inches(0.5),
                "Read the compendium  \u2022  run the demos  \u2022  the source is on the repo",
                size=F.BODY, color=C.TEXT_MUTED, font=FontFam.BODY)
    bp.add_text(slide, Inches(1.0), Inches(6.4), Inches(11.3), Inches(0.4),
                REPO_URL, size=F.CAPTION, color=C.TEXT_MUTED, font=FontFam.CODE)
    bp.set_notes(slide,
        "To close: statelessness isn't a checkbox you tick in code — it's a "
        "posture the orchestrator can rely on, and the C++ techniques in this "
        "compendium are how a C++ service earns it. Request scope through RAII "
        "and PMR, process scope owned explicitly in main(), authoritative "
        "state externalized so any replica can die without loss.\n\n"
        "And none of it is theory: every pattern has a runnable demo, and the "
        "whole arc from RAII to the gRPC capstone is host-verified. Clone the "
        "repo, run the demos, read the long-form compendium. Thank you.")


def main():
    sys.path.insert(0, str(TOOLS))
    from sections_statelessness import SECTIONS

    total = 2  # title + agenda
    for s in SECTIONS:
        total += 1 + len(s["slides"])
    total += 1  # closing
    print(f"Total slides planned: {total}")

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    page = 1
    print(f"  [{page:3d}/{total}] Title"); build_title(prs, total); page += 1
    print(f"  [{page:3d}/{total}] Agenda"); build_agenda(prs, total, SECTIONS); page += 1

    for section in SECTIONS:
        print(f"  [{page:3d}/{total}] \u00a7{section['num']} divider \u2014 {section['title']}")
        bp.build_section_divider(prs, section["num"], section["title"],
                                 section["tagline"], page, total,
                                 section["divider_notes"])
        page += 1
        for slide_data in section["slides"]:
            kind = slide_data["kind"]
            t = slide_data.get("title") or slide_data.get("demo_name", "(demo)")
            print(f"  [{page:3d}/{total}] \u00a7{section['num']} {kind} \u2014 {t[:48]}")
            bp.render_slide(prs, section, slide_data, page, total)
            page += 1

    print(f"  [{page:3d}/{total}] Closing"); build_closing(prs, total); page += 1

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT_PATH))
    print(f"\nSaved: {OUT_PATH}")
    print(f"Slide count: {page - 1}")
    NOTES_PATH.write_text(
        f"# Build notes — cpp-statelessness-compendium.pptx\n\n"
        f"- Total slides: {page - 1}\n"
        f"- Sections rendered: {len(SECTIONS)}\n"
        f"- Generator: tools/build-statelessness-pptx.py\n"
        f"- Content source: tools/sections_statelessness.py\n"
        f"- Shared builders: tools/build-pptx.py (loaded via importlib)\n")
    print(f"Wrote: {NOTES_PATH}")


if __name__ == "__main__":
    main()
